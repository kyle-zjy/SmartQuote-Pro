"""
Regenerates src/data/colours.json from the Goldco standard colour workbook.

    python scripts/generate-colours.py
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "documents" / "Goldco Standard Colour List (Updated Dec-26).xlsx"
OUT_PATH = ROOT / "src" / "data" / "colours.json"

# Spreadsheet columns -> product keys used in pricing.json
COLUMN_PRODUCTS = {
    3: ["supascreen"],
    4: ["supascreen"],
    5: ["supascreen"],
    6: ["intrudaguard"],
    7: ["intrudaguard"],
    8: ["intrudaguard"],
    9: ["flyscreens", "7mm-diamond"],
    10: ["flyscreens"],
    11: ["flyscreens"],
    12: ["7mm-diamond"],
    13: ["7mm-diamond"],
}

ALIASES = [
    {"name": "White", "mapsTo": "White Birch Gloss"},
    {"name": "Black", "mapsTo": "Black Satin"},
    {"name": "Bronze", "mapsTo": "10um T37 Bronze"},
]


def clean_name(raw: str) -> str:
    return re.sub(r"\s+", " ", raw.replace("\n", " ")).strip()


def display_name(raw: str) -> str:
    return clean_name(re.split(r"\*\*\*", raw, maxsplit=1)[0]).strip(" -")


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Standard Colours List"]
    colours = []
    for row in range(3, ws.max_row + 1):
        raw = ws.cell(row=row, column=2).value
        if not isinstance(raw, str) or not raw.strip():
            continue
        products: set[str] = set()
        for col, keys in COLUMN_PRODUCTS.items():
            if str(ws.cell(row=row, column=col).value).strip().lower() == "yes":
                products.update(keys)
        colours.append(
            {
                "name": display_name(raw),
                "additionalCharge": "additional charges apply" in raw.lower(),
                "products": sorted(products),
            }
        )

    by_name = {c["name"]: c for c in colours}
    for alias in ALIASES:
        source = next((c for c in colours if c["name"].startswith(alias["mapsTo"])), None)
        if source and alias["name"] not in by_name:
            colours.insert(
                0,
                {
                    "name": alias["name"],
                    "additionalCharge": source["additionalCharge"],
                    "products": source["products"],
                },
            )

    colours.append({"name": "Non-standard / Other", "additionalCharge": True, "products": []})
    OUT_PATH.write_text(json.dumps(colours, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {len(colours)} colours to {OUT_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
