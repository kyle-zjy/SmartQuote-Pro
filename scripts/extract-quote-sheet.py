"""
Extracts the student quote-sheet workbook into JSON + config drawings.

    python scripts/extract-quote-sheet.py

Source: documents/AAA - Quote sheet - Student Version WIP.xlsx
"""
import json
import shutil
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "documents" / "AAA - Quote sheet - Student Version WIP.xlsx"
OUT_JSON = ROOT / "src" / "data" / "quoteSheet.json"
OUT_IMAGES = ROOT / "src" / "assets" / "sheet-codes"

HEIGHT_COLS = list(range(7, 10))  # G-I = H1-H3
WIDTH_COLS = list(range(10, 20))  # J-S = W1-W10


def yes(value) -> bool:
    return isinstance(value, str) and value.strip().lower() == "yes"


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    images = wb["Images"]
    data = wb["Data Sheet"]
    quote = wb["Quote Sheet"]

    products = [cell.value.strip() for cell in data["B"] if isinstance(cell.value, str) and cell.value.strip() != "Product"]
    screen_types = [cell.value.strip() for cell in data["D"] if isinstance(cell.value, str) and cell.value.strip() != "Screen type"]

    notes = []
    for row in range(13, 27):
        text = quote.cell(row=row, column=6).value
        if isinstance(text, str) and text.strip():
            notes.append(text.strip())

    configs = []
    for r in range(2, images.max_row + 1):
        code = images.cell(row=r, column=1).value
        if not isinstance(code, str) or not code.strip():
            continue
        code = code.strip()
        configs.append({
            "code": code,
            "image": f"{code.lower()}.png",
            "panels": images.cell(row=r, column=3).value or 1,
            "widthFactor": images.cell(row=r, column=4).value or 1,
            "widthOffset": images.cell(row=r, column=5).value or 0,
            "heightFactor": images.cell(row=r, column=6).value or 1,
            "heightPoints": [f"H{i}" for i, col in enumerate(HEIGHT_COLS, start=1) if yes(images.cell(row=r, column=col).value)],
            "widthPoints": [f"W{i}" for i, col in enumerate(WIDTH_COLS, start=1) if yes(images.cell(row=r, column=col).value)],
        })

    OUT_JSON.write_text(
        json.dumps(
            {
                "products": products,
                "screenTypes": screen_types,
                "notes": notes,
                "configs": configs,
            },
            indent=2,
        ) + "\n",
        encoding="utf-8",
    )

    OUT_IMAGES.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(XLSX_PATH) as zipped:
        for index, config in enumerate(configs, start=1):
            dest = OUT_IMAGES / config["image"]
            dest.write_bytes(zipped.read(f"xl/media/image{index}.png"))

    print(f"Wrote {OUT_JSON} ({len(configs)} configs)")
    print(f"Wrote {len(configs)} drawings into {OUT_IMAGES}")


if __name__ == "__main__":
    main()
