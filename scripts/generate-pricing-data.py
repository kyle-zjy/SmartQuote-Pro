"""
Regenerates src/data/pricing.json and src/data/addons.json from the master
price list spreadsheet in documents/.

Run whenever the xlsx is updated:

    python scripts/generate-pricing-data.py

Requires: pip install openpyxl
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "documents" / "Supply & Install (17-5-26).xlsx"
OUT_PRICING = ROOT / "src" / "data" / "pricing.json"
OUT_ADDONS = ROOT / "src" / "data" / "addons.json"

MAX_COL = 30


def read_row(ws, row, max_col=MAX_COL):
    return [ws.cell(row=row, column=c).value for c in range(1, max_col + 1)]


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def find_title_row(ws, substring, start_row=1, col_range=(1, MAX_COL)):
    """First row where some cell (case-insensitive) contains `substring`."""
    needle = substring.lower()
    for r in range(start_row, ws.max_row + 1):
        for c in range(*col_range):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and needle in v.lower():
                return r
    raise ValueError(f"title containing {substring!r} not found in sheet {ws.title!r}")


def numeric_run(row_values):
    """Index + values of the first contiguous run of numbers in a row."""
    start = None
    vals = []
    for i, v in enumerate(row_values):
        if is_number(v):
            if start is None:
                start = i
            vals.append(v)
        elif start is not None:
            break
    return start, vals


def extract_matrix(ws, title_substring):
    """A `<Title>` row followed directly by a width header row, then one
    row per height with prices across. Stops at the first row whose height
    cell isn't numeric. Cell value 'N/A' becomes None."""
    title_row = find_title_row(ws, title_substring)
    header_row = title_row + 1
    width_col_idx, widths = numeric_run(read_row(ws, header_row))
    if width_col_idx is None:
        raise ValueError(f"no width header found under {title_substring!r} in {ws.title!r}")
    height_col_idx = width_col_idx - 1

    heights = []
    prices = []
    r = header_row + 1
    while True:
        row_vals = read_row(ws, r)
        height = row_vals[height_col_idx] if height_col_idx >= 0 else None
        if not is_number(height):
            break
        row_prices = row_vals[width_col_idx: width_col_idx + len(widths)]
        prices.append([round(p) if is_number(p) else None for p in row_prices])
        heights.append(height)
        r += 1

    return {"widths": widths, "heights": heights, "prices": prices}


def extract_extras(ws, title_substring):
    """A `<X> EXTRAS` row with two labelled columns ('UNDER n HIGH ADD',
    'OVER n HIGH ADD') a few cells to the right, then one row per mesh
    option in the same label column as the title."""
    title_row = find_title_row(ws, title_substring)
    row_vals = read_row(ws, title_row)
    label_col_idx = next(i for i, v in enumerate(row_vals) if isinstance(v, str) and title_substring.lower() in v.lower())

    under_col_idx = under_threshold = None
    over_col_idx = over_threshold = None
    for i, v in enumerate(row_vals):
        if not isinstance(v, str):
            continue
        m = re.search(r"UNDER\s+(\d+)\s+HIGH", v, re.I)
        if m:
            under_col_idx, under_threshold = i, int(m.group(1))
        m = re.search(r"OVER\s+(\d+)\s+HIGH", v, re.I)
        if m:
            over_col_idx, over_threshold = i, int(m.group(1))

    options = []
    r = title_row + 1
    while True:
        vals = read_row(ws, r)
        name = vals[label_col_idx]
        if not isinstance(name, str) or not name.strip():
            break
        under_val = vals[under_col_idx] if under_col_idx is not None else None
        over_val = vals[over_col_idx] if over_col_idx is not None else None
        options.append({
            "name": name.strip(),
            "under": round(under_val) if is_number(under_val) else under_val,
            "over": round(over_val) if is_number(over_val) else over_val,
        })
        r += 1

    threshold_mm = under_threshold or over_threshold
    return {"thresholdMm": threshold_mm, "options": options}


def extract_pricing_as_at(ws):
    for row in ws.iter_rows(max_row=5):
        for cell in row:
            if isinstance(cell.value, str) and "pricing as at" in cell.value.lower():
                date_cell = ws.cell(row=cell.row, column=cell.column + 1)
                v = date_cell.value
                return v.date().isoformat() if hasattr(v, "date") else v
    return None


def build_products(wb):
    products = []

    # Supascreen
    ws = wb["Supascreen"]
    products.append({
        "key": "supascreen",
        "name": "Supascreen",
        "pricingAsAt": extract_pricing_as_at(ws),
        "categories": [
            {"key": "windows", "label": "Windows", **extract_matrix(ws, "Supascreen  Windows"), "extras": None},
            {"key": "doors", "label": "Doors", **extract_matrix(ws, "Supascreen  Doors"), "extras": None},
        ],
    })

    # IntrudaGuard (sheet name has a leading space)
    ws = wb[" IntrudaGuard"]
    products.append({
        "key": "intrudaguard",
        "name": "IntrudaGuard",
        "pricingAsAt": extract_pricing_as_at(ws),
        "categories": [
            {"key": "windows", "label": "Windows", **extract_matrix(ws, "IntrudaGuard Windows"), "extras": None},
            {"key": "doors", "label": "Doors", **extract_matrix(ws, "IntrudaGuard Doors"), "extras": None},
        ],
    })

    # 7mm Diamond
    ws = wb["7mm Diamond"]
    products.append({
        "key": "7mm-diamond",
        "name": "7mm Diamond",
        "pricingAsAt": extract_pricing_as_at(ws),
        "categories": [
            {"key": "windows", "label": "Windows", **extract_matrix(ws, "7mm Diamond  Windows"),
             "extras": extract_extras(ws, "WINDOWS EXTRAS")},
            {"key": "doors", "label": "Doors", **extract_matrix(ws, "7mm Diamond Doors"),
             "extras": extract_extras(ws, "DOOR EXTRAS")},
        ],
    })

    # Fly Screens
    ws = wb["Fly Screens"]
    products.append({
        "key": "flyscreens",
        "name": "Fly Screens",
        "pricingAsAt": extract_pricing_as_at(ws),
        "note": "Additional charge of $15+GST for double hung windows",
        "categories": [
            {"key": "windows", "label": "Windows (Standard Mesh)", **extract_matrix(ws, "Flyscreens - Standard Mesh"),
             "extras": extract_extras(ws, "WINDOWS EXTRAS")},
            {"key": "sliding-doors", "label": "Sliding Doors (Standard Mesh)",
             **extract_matrix(ws, "Flyscreen Sliding Doors"), "extras": None},
            {"key": "hinged-doors", "label": "Hinged Doors (Standard Mesh)",
             **extract_matrix(ws, "Flyscreen Hinged Doors"), "extras": extract_extras(ws, "DOOR EXTRAS")},
        ],
    })

    return products


def build_addons(wb):
    ws = wb["Retail Supply Extras"]
    addons = []
    section = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        vals = [c.value for c in row]
        label = vals[1] if len(vals) > 1 else None
        if not isinstance(label, str) or not label.strip():
            continue
        if label.strip().upper() in ("ADDONS", "EXTRAS"):
            section = label.strip().title()
            continue
        if label.strip().upper().startswith(("PRODUCT LIST", "ALL PRICES", "EFFECTIVE DATE", "RETAIL")):
            continue
        if section is None:
            continue  # skip anything before the first ADDONS/EXTRAS section header
        price = vals[2] if len(vals) > 2 else None
        unit = vals[3] if len(vals) > 3 else None
        addons.append({
            "section": section,
            "name": label.strip(),
            "price": round(price) if is_number(price) else None,
            "priceOnRequest": isinstance(price, str) and "request" in price.lower(),
            "unit": unit if isinstance(unit, str) else None,
        })
    return addons


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    pricing = {"products": build_products(wb), "note": "All prices exclude GST"}
    addons = build_addons(wb)

    OUT_PRICING.parent.mkdir(parents=True, exist_ok=True)
    OUT_PRICING.write_text(json.dumps(pricing, indent=2), encoding="utf-8")
    OUT_ADDONS.write_text(json.dumps(addons, indent=2), encoding="utf-8")
    print(f"Wrote {OUT_PRICING}")
    print(f"Wrote {OUT_ADDONS}")


if __name__ == "__main__":
    main()
